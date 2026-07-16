# -*- coding: utf-8 -*-
"""v2 确定性回归：canonical handoff、引导语治理和导出门禁。"""
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


count_chars = load_module("count_chars_v2", SCRIPTS / "count_chars.py")
pick_guidance = load_module("pick_guidance_v2", SCRIPTS / "pick_guidance.py")
validate_output = load_module("validate_output_v2", SCRIPTS / "validate_output.py")
build_docx = load_module("build_docx_v2", SCRIPTS / "build_docx.py")
handoff_contract = load_module("handoff_contract_v3", SCRIPTS / "handoff_contract.py")


def lock(raw, locked_text=None, normalization=None, provenance="source_verbatim"):
    return {
        "raw": raw,
        "locked_text": locked_text or raw,
        "normalization": normalization or [],
        "provenance": provenance,
        "requires_approval": False,
    }


def valid_item(script=None, source_text=None):
    if script is None:
        script = "腰腿疼痛遇冷容易反复，可以准备桃树叶一大把、生姜五片（需辨证为寒湿痹阻）记不住先收藏，用的时候翻出来看。"
    if source_text is None:
        source_text = "腰腿疼痛反复发作，取桃树叶一大把、生姜五片泡脚。"
    return {
        "source": {"link": "https://example.com/video", "title": "腰腿疼痛泡脚方", "text": source_text},
        "locks": {
            "formula_items": [lock("桃树叶一大把"), lock("生姜五片")],
            "diseases": [lock("腰腿疼痛")],
            "syndromes": [lock("寒湿痹阻", provenance="review_approved")],
            "numeric_hooks": [],
        },
        "title_decision": {"type": "unchanged", "rationale": ""},
        "manual_checks": {key: True for key in build_docx.MANUAL_CHECK_KEYS},
        "warning_decisions": {},
        "public": {
            "link": "https://example.com/video",
            "title": "腰腿疼痛泡脚方",
            "script": script,
            "notes": [],
            "tips": [],
            "processing": [],
        },
    }


def valid_handoff(items=None, filename="腰腿疼痛-洗稿文档"):
    return {"schema_version": "2", "stage": "rewritten", "filename": filename, "items": items or [valid_item()]}


class HandoffContractV3Tests(unittest.TestCase):
    def test_v2_migration_preserves_source_backed_locks(self):
        old = {
            "schema_version": "2",
            "stage": "parsed_checked",
            "source": {"link": "", "title": "测试标题", "text": "方中用荷叶10g，辨证为湿热。"},
            "internal_review": {"licensed_physician_review_required": True},
            "locks": {
                "formula_items": [{
                    "raw": "荷叶10g", "locked_text": "荷叶10克",
                    "normalization": ["unit:g→克"], "provenance": "source_normalized",
                    "requires_approval": False,
                }],
                "diseases": [], "syndromes": [], "numeric_hooks": [],
            },
        }
        migrated = handoff_contract.v2_to_v3(old)
        hook = migrated["hooks"]["verbatim"][0]
        self.assertEqual("approved", hook["approval_status"])
        self.assertEqual("verbatim", hook["safety_disposition"])
        self.assertEqual("荷叶10克", hook["delivery_text"])

    def test_v2_ambiguous_or_review_approved_locks_become_pending(self):
        old = {
            "schema_version": "2", "stage": "parsed_checked",
            "source": {"link": "", "title": "测试", "text": "原文没有新增证型。"},
            "internal_review": {"licensed_physician_review_required": True},
            "locks": {
                "formula_items": [], "diseases": [],
                "syndromes": [{
                    "raw": "肝胆湿热", "locked_text": "肝胆湿热", "normalization": [],
                    "provenance": "review_approved", "requires_approval": False,
                }],
                "numeric_hooks": [],
            },
        }
        migrated = handoff_contract.v2_to_v3(old)
        hook = migrated["hooks"]["verbatim"][0]
        self.assertEqual("pending", hook["approval_status"])
        self.assertEqual("pending", hook["safety_disposition"])

    def test_unapproved_hook_cannot_be_verbatim(self):
        source = "方中用荷叶10克。"
        state = {
            "schema_version": "3", "stage": "parsed_checked",
            "source": {"link": "", "title": "测试", "text": source},
            "hooks": {"verbatim": [{
                "id": "formula-1", "kind": "formula_item", "raw": "荷叶10克",
                "delivery_text": "荷叶10克", "source_span": [3, 8], "normalization_ops": [],
                "provenance": "source_verbatim", "approval_status": "pending",
                "safety_disposition": "verbatim",
            }], "semantic_strength": []},
            "review_findings": [],
            "title_review": {"decision": "unchanged", "title": "测试", "reason_codes": [], "rationale": ""},
            "internal_review": {"licensed_physician_review_required": True},
        }
        with self.assertRaisesRegex(ValueError, "未批准"):
            handoff_contract.validate_parsed_state(state)

    def test_v2_pure_string_and_unverifiable_normalization_become_pending(self):
        pure_string = {
            "schema_version": "2", "stage": "parsed_checked",
            "source": {"link": "", "title": "测试", "text": "原文声称根治。"},
            "internal_review": {"licensed_physician_review_required": True},
            "locks": {"formula_items": [], "diseases": [], "syndromes": [], "numeric_hooks": ["根治"]},
        }
        migrated = handoff_contract.v2_to_v3(pure_string)
        self.assertEqual("pending", migrated["hooks"]["verbatim"][0]["approval_status"])

        invalid_normalization = {
            "schema_version": "2", "stage": "parsed_checked",
            "source": {"link": "", "title": "测试", "text": "方中用荷叶10g。"},
            "internal_review": {"licensed_physician_review_required": True},
            "locks": {
                "formula_items": [{
                    "raw": "荷叶10g", "locked_text": "砒霜100克",
                    "normalization": ["unit:g→克"], "provenance": "source_normalized",
                    "requires_approval": False,
                }],
                "diseases": [], "syndromes": [], "numeric_hooks": [],
            },
        }
        migrated = handoff_contract.v2_to_v3(invalid_normalization)
        self.assertEqual("pending", migrated["hooks"]["verbatim"][0]["approval_status"])

    def test_unverified_semantic_identity_cannot_be_verbatim(self):
        source = "我是三甲名医，今天分享方法。"
        state = {
            "schema_version": "3", "stage": "parsed_checked",
            "source": {"link": "", "title": "测试", "text": source},
            "hooks": {"verbatim": [], "semantic_strength": [{
                "id": "identity-1", "kind": "identity", "source_text": "我是三甲名医",
                "source_span": [0, 6], "truth_status": "unverified", "safety_disposition": "verbatim",
            }]},
            "review_findings": [],
            "title_review": {"decision": "unchanged", "title": "测试", "reason_codes": [], "rationale": ""},
            "internal_review": {"licensed_physician_review_required": True},
        }
        with self.assertRaisesRegex(ValueError, "未核实|unverified"):
            handoff_contract.validate_parsed_state(state)

    def test_review_approved_non_contiguous_hook_uses_sentinel_span(self):
        source = "方中提到青葙子，取10到15克。"
        hook = {
            "id": "formula-1", "kind": "formula_item", "raw": "青葙子10到15克",
            "delivery_text": "青葙子10到15克", "source_span": [-1, -1],
            "normalization_ops": ["combine:药材与剂量合并保全"],
            "provenance": "review_approved", "approval_status": "approved",
            "safety_disposition": "verbatim",
        }
        handoff_contract.validate_verbatim_hook(hook, source, "hook")
        hook["source_span"] = [0, 2]
        with self.assertRaisesRegex(ValueError, "source_span"):
            handoff_contract.validate_verbatim_hook(hook, source, "hook")

    def test_engagement_hook_is_valid_and_function_is_preserved(self):
        source = "愿意的话送朵小花，留句话。"
        hooks = {"verbatim": [], "semantic_strength": [{
            "id": "engagement-1", "kind": "engagement", "source_text": "送朵小花，留句话",
            "source_span": [4, 12], "truth_status": "verified", "safety_disposition": "demoted",
        }]}
        handoff_contract.validate_hooks(hooks, source)
        handoff_contract.validate_engagement_preservation(hooks, "觉得实用就点个红心。")
        with self.assertRaisesRegex(ValueError, "互动功能未保留"):
            handoff_contract.validate_engagement_preservation(hooks, "记住这个方法。")

    def test_title_reason_codes_and_number_preservation(self):
        source_title = "90%的胆结石，记好一个排石方"
        handoff_contract.validate_title_review(
            {"decision": "unchanged", "title": source_title, "reason_codes": [], "rationale": ""},
            source_title,
        )
        with self.assertRaisesRegex(ValueError, "非硬红线理由"):
            handoff_contract.validate_title_review(
                {"decision": "safety_adjusted", "title": "查出胆结石，记好这个思路", "reason_codes": ["unsupported_percentage"], "rationale": "删比例"},
                source_title,
            )
        with self.assertRaisesRegex(ValueError, "不得删除原标题数字"):
            handoff_contract.validate_title_review(
                {"decision": "safety_adjusted", "title": "胆结石，记好一个排石方", "reason_codes": ["efficacy_guarantee"], "rationale": "处理危险保证"},
                source_title,
            )


class CountCharsV2Tests(unittest.TestCase):
    def test_boundaries_and_newlines(self):
        self.assertEqual(300, count_chars.count("字" * 150 + "\r\n" + "字" * 150))
        self.assertEqual(301, count_chars.count("字" * 301))


class ValidateOutputV2Tests(unittest.TestCase):
    def test_empty_script_fails(self):
        self.assertIn("empty-script", {x["id"] for x in validate_output.validate_text("")["errors"]})

    def test_300_passes_and_301_fails(self):
        self.assertFalse(validate_output.validate_text("字" * 300)["errors"])
        self.assertIn("max-chars", {x["id"] for x in validate_output.validate_text("字" * 301)["errors"]})

    def test_forbidden_punctuation_and_units_fail(self):
        for text in ("注意：慢慢看", "这件事——要记住", "他说“可以”", '写成"这样"', "荷叶10g"):
            with self.subTest(text=text):
                self.assertTrue(validate_output.validate_text(text)["errors"])

    def test_all_lock_types_pass_when_preserved(self):
        locks = {
            "formula_items": [lock("桃树叶一大把"), lock("生姜五片")],
            "diseases": [lock("腰腿疼痛")],
            "syndromes": [lock("寒湿痹阻")],
            "numeric_hooks": [lock("90%")],
        }
        text = "90%的腰腿疼痛可以准备桃树叶一大把、生姜五片（需辨证为寒湿痹阻）"
        self.assertFalse(validate_output.validate_text(text, locks)["errors"])

    def test_each_missing_lock_fails(self):
        for key, value in (("formula_items", "荷叶10克"), ("diseases", "脂肪肝"), ("syndromes", "痰湿内盛"), ("numeric_hooks", "90%")):
            locks = {name: [] for name in build_docx.LOCK_KEYS}
            locks[key] = [lock(value)]
            with self.subTest(key=key):
                issue_ids = {x["id"] for x in validate_output.validate_text("普通正文", locks)["errors"]}
                expected_id = "missing-syndrome-annotation" if key == "syndromes" else "missing-lock"
                self.assertIn(expected_id, issue_ids)

    def test_medical_claims_and_self_treatment_block(self):
        for text in ("这个病的根子就是湿热。", "这些表现全都是气血不足引起的。", "小毛病自己就能调。", "保证你以后不用跑医院。"):
            with self.subTest(text=text):
                self.assertTrue(validate_output.validate_text(text)["errors"])


    def test_editorial_tone_blocks_known_failures_and_variants(self):
        bad_cases = (
            "关键是先看自己到底适不适合。",
            "重点是要先判断这个方法是否适合自己。",
            "视物模糊不一定都是这个原因。",
            "这些表现未必全由湿热导致。",
            "在辨证合适的前提下再使用。",
            "先把原因和证型辨清，再看这个思路适不适合。",
        )
        for text in bad_cases:
            with self.subTest(text=text):
                self.assertIn("editorial-tone", {x["id"] for x in validate_output.validate_text(text)["errors"]})

    def test_editorial_tone_does_not_block_natural_oral_phrasing(self):
        good_cases = (
            "平时腰膝酸软的朋友可以试试。",
            "中医讲究辨证论治。",
            "在中医看来，这类情况常和湿热有关。",
            "这类视物模糊常和用眼疲劳有关。",
            "带状疱疹72小时内需及时就医，外敷只是辅助。",
            "突然胸痛伴大汗或呼吸困难，要立即就医。",
            "症状突然加重，尽快去急诊。",
        )
        for text in good_cases:
            with self.subTest(text=text):
                self.assertNotIn("editorial-tone", {x["id"] for x in validate_output.validate_text(text)["errors"]})

    def test_syndrome_annotation_authorization_format_and_count(self):
        no_syndrome = {name: [] for name in build_docx.LOCK_KEYS}
        self.assertFalse(validate_output.validate_text("普通口播正文", no_syndrome)["errors"])
        self.assertIn(
            "unexpected-syndrome-annotation",
            {x["id"] for x in validate_output.validate_text("普通正文（需辨证为肝火上炎）", no_syndrome)["errors"]},
        )

        locks = {name: [] for name in build_docx.LOCK_KEYS}
        locks["syndromes"] = [lock("肝火上炎")]
        good = "能帮助清肝明目（需辨证为肝火上炎）"
        self.assertFalse(validate_output.validate_text(good, locks)["errors"])

        bad_cases = {
            "肝火上炎的人可以了解这个思路。": "missing-syndrome-annotation",
            "（需辨证为肝火上炎）（需辨证为肝火上炎）": "duplicate-syndrome-annotation",
            "（需辨证为痰湿内盛）": "invalid-syndrome-annotation",
            "（需辨证为肝火上炎，伴眼红眼干）": "invalid-syndrome-annotation",
            "(需辨证为肝火上炎)": "invalid-syndrome-annotation",
            "（辨证为肝火上炎）": "invalid-syndrome-annotation",
            "需辨证为肝火上炎。": "invalid-syndrome-annotation",
            "在辨证合适的前提下，能帮助清肝明目（需辨证为肝火上炎）": "editorial-tone",
        }
        for text, issue_id in bad_cases.items():
            with self.subTest(text=text):
                self.assertIn(issue_id, {x["id"] for x in validate_output.validate_text(text, locks)["errors"]})

    def test_formula_lock_accepts_natural_measure_word_context(self):
        locks = {"formula_items": [lock("荷叶10克")], "diseases": [], "syndromes": [], "numeric_hooks": []}
        issue_ids = {x["id"] for x in validate_output.validate_text("这味荷叶10克可以了解。", locks)["errors"]}
        self.assertNotIn("missing-lock", issue_ids)

    def test_marketing_phrase_cannot_abuse_meta_language_negation(self):
        text = "不要说我没提醒你，这个方法保证治好。"
        issue_ids = {x["id"] for x in validate_output.validate_text(text)["errors"]}
        self.assertIn("medical-guarantee", issue_ids)

    def test_ten_character_copy_and_lock_exemption(self):
        source = "今天这个珍藏多年的方法一定要认真看完"
        self.assertIn("verbatim-copy", {x["id"] for x in validate_output.validate_text(source, source_text=source)["errors"]})
        locks = {"formula_items": [], "diseases": [], "syndromes": [], "numeric_hooks": [lock(source)]}
        self.assertNotIn("verbatim-copy", {x["id"] for x in validate_output.validate_text(source, locks, source)["errors"]})

    def test_repeated_today_preview_is_blocked(self):
        bad_text = "今天要说的方法很简单。今天跟大家分享的这招，很多人花钱都学不到。"
        issue_ids = {x["id"] for x in validate_output.validate_text(bad_text)["errors"]}
        self.assertIn("today-preview-repeat", issue_ids)

    def test_single_today_preview_passes(self):
        good_text = "今天要说的方法很简单，青葙子10到15克煎水服用，记不住先收藏。"
        self.assertEqual([], validate_output.validate_text(good_text)["errors"])

    def test_defensive_popsci_templates_are_blocked(self):
        bad_cases = (
            "这是常见的配伍思路。",
            "具体能不能用，要让中医师结合体质判断。",
            "是否需要手术，应由专科医生综合评估。",
            "这个方子不能替代规范诊疗。",
            "建议咨询专业医生。",
        )
        for text in bad_cases:
            with self.subTest(text=text):
                issue_ids = {x["id"] for x in validate_output.validate_text(text)["errors"]}
                self.assertTrue({"popsci-bridge", "editorial-tone"}.intersection(issue_ids), text)

    def test_popsci_bridge_transitions_are_blocked(self):
        for text in ("在中医看来和肝火有关，常会用到青葙子", "临床上常用这味药", "顺着这个思路，一般会用"):
            with self.subTest(text=text):
                issue_ids = {x["id"] for x in validate_output.validate_text(text)["errors"]}
                self.assertIn("popsci-bridge", issue_ids)

    def test_confirmed_medical_safety_matrix(self):
        unsafe = (
            "这个方法保证治好。",
            "这样做一定能治愈。",
            "不用去医院，在家自己治。",
            "这个方子能治疗脂肪肝。",
            "确保以后不会复发。",
        )
        for text in unsafe:
            with self.subTest(text=text):
                self.assertTrue(validate_output.validate_text(text)["errors"], text)

        safe_negations = (
            "这个方法不能保证治好。",
            "不要写保证治好这样的说法。",
            "并非不用去医院，有不适应及时就医。",
        )
        for text in safe_negations:
            with self.subTest(text=text):
                self.assertEqual([], validate_output.validate_text(text)["errors"], text)

    def test_lock_match_requires_entity_boundary(self):
        locks = {"formula_items": [lock("荷叶10克")], "diseases": [], "syndromes": [], "numeric_hooks": []}
        issue_ids = {x["id"] for x in validate_output.validate_text("薄荷叶10克可以使用。", locks)["errors"]}
        self.assertIn("missing-lock", issue_ids)

    def test_syndrome_marker_outside_approved_annotation_is_blocked(self):
        locks = {name: [] for name in build_docx.LOCK_KEYS}
        locks["syndromes"] = [lock("肝火上炎")]
        text = "先需辨证为痰湿内盛，再看（需辨证为肝火上炎）"
        issue_ids = {x["id"] for x in validate_output.validate_text(text, locks)["errors"]}
        self.assertIn("invalid-syndrome-annotation", issue_ids)

    def test_lock_does_not_split_verbatim_copy_window(self):
        source = "脂肪肝值得每个人认真收藏这段内容"
        locks = {"formula_items": [], "diseases": [lock("脂肪肝")], "syndromes": [], "numeric_hooks": []}
        issue_ids = {x["id"] for x in validate_output.validate_text(source, locks, source)["errors"]}
        self.assertIn("verbatim-copy", issue_ids)


class GuidanceLibraryV2Tests(unittest.TestCase):
    def test_library_schema_and_status(self):
        items = pick_guidance.load_items(ROOT / "references" / "guidance_library.json")
        self.assertGreater(len(items), 20)
        for item in items:
            self.assertEqual(item["char_count"], len(item["text"]))
            self.assertIn(item["status"], {"approved", "quarantined"})
            self.assertNotIn("您", item["text"])

    def test_known_risky_entries_are_quarantined(self):
        items = pick_guidance.load_items(ROOT / "references" / "guidance_library.json")
        risky = [x for x in items if any(word in x["text"] for word in ("小毛病，咱不求人", "药店就要干不下去", "回去试一试", "腰腿疼痛"))]
        self.assertTrue(risky)
        self.assertTrue(all(x["status"] == "quarantined" for x in risky))

    def test_seed_is_reproducible_and_budget_is_respected(self):
        items = [x for x in pick_guidance.load_items(ROOT / "references" / "guidance_library.json") if x["status"] == "approved" and x["char_count"] <= 75]
        self.assertEqual(pick_guidance.pick(items, 3, seed=42), pick_guidance.pick(items, 3, seed=42))
        self.assertTrue(all(x["char_count"] <= 75 for x in pick_guidance.pick(items, 3, seed=42)))

    def test_invalid_object_schema_fails(self):
        with self.assertRaises(pick_guidance.LibraryError):
            pick_guidance.parse_entry({"text": "内容", "universal": True}, 1)
    def test_approved_guidance_is_validator_compatible(self):
        items = pick_guidance.load_items(ROOT / "references" / "guidance_library.json")
        for item in (x for x in items if x["status"] == "approved"):
            with self.subTest(item=item["id"]):
                report = validate_output.validate_text(item["text"])
                self.assertEqual([], report["errors"], item["id"])
                self.assertEqual([], report["warnings"], item["id"])

    def test_guidance_metadata_prevents_wrong_placement_and_identity_injection(self):
        items = pick_guidance.load_items(ROOT / "references" / "guidance_library.json")
        by_id = {item["id"]: item for item in items}
        self.assertEqual("closing", by_id["guidance-036"]["placement"])
        self.assertNotIn("general_health", by_id["guidance-040"]["content_types"])
        for item_id in ("guidance-006", "guidance-008", "guidance-027"):
            self.assertNotEqual("approved", by_id[item_id]["review_status"])
            self.assertTrue(by_id[item_id]["persona_requirements"])

    def test_high_similarity_status_conflicts_are_rejected(self):
        items = pick_guidance.load_items(ROOT / "references" / "guidance_library.json")
        by_id = {item["id"]: item for item in items}
        self.assertEqual(by_id["guidance-031"]["review_status"], by_id["guidance-034"]["review_status"])


class BuildDocxV2Tests(unittest.TestCase):
    def test_valid_handoff_exports_exact_validated_script(self):
        data = valid_handoff()
        with tempfile.TemporaryDirectory() as output_dir:
            path, count, reports = build_docx.build(data, output_dir)
            text = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertEqual(1, count)
        self.assertFalse(reports[0]["errors"])
        self.assertIn(data["items"][0]["public"]["script"], text)

    def test_invalid_script_does_not_export(self):
        item = valid_item(script="")
        with tempfile.TemporaryDirectory() as output_dir:
            with self.assertRaises(ValueError):
                build_docx.build(valid_handoff([item]), output_dir)
            self.assertEqual([], list(Path(output_dir).iterdir()))

    def test_missing_lock_manifest_fails(self):
        data = valid_handoff()
        del data["items"][0]["locks"]["numeric_hooks"]
        with self.assertRaises(ValueError):
            build_docx.validate_handoff(data)

    def test_unhandled_warning_blocks_export(self):
        item = valid_item(script="腰腿疼痛可以准备桃树叶一大把、生姜五片，特别实用（需辨证为寒湿痹阻）")
        with self.assertRaisesRegex(ValueError, "未处置"):
            build_docx.validate_handoff(valid_handoff([item]))
        item["warning_decisions"] = {"degree-word": {"status": "false_positive", "reason": "原文即为特别实用，未升级"}}
        build_docx.validate_handoff(valid_handoff([item]))

    def test_resolved_cannot_bypass_warning_that_still_exists(self):
        item = valid_item(script="腰腿疼痛可以准备桃树叶一大把、生姜五片，特别实用（需辨证为寒湿痹阻）")
        item["warning_decisions"] = {"degree-word": "resolved"}
        with self.assertRaisesRegex(ValueError, "仍存在|无效|理由"):
            build_docx.validate_handoff(valid_handoff([item]))

    def test_title_and_public_sections_are_medically_scanned(self):
        title_item = valid_item()
        title_item["source"]["title"] = "不用治疗，保证治好"
        title_item["public"]["title"] = "不用治疗，保证治好"
        with self.assertRaisesRegex(ValueError, "标题|医疗|保证|治疗"):
            build_docx.validate_handoff(valid_handoff([title_item]))

        tips_item = valid_item()
        tips_item["public"]["tips"] = [{"text": "无需就医，这个方法保证治好。"}]
        with self.assertRaisesRegex(ValueError, "tips|医疗|保证|就医"):
            build_docx.validate_handoff(valid_handoff([tips_item]))

    def test_safety_adjusted_title_must_change_and_be_safe(self):
        item = valid_item()
        item["title_decision"] = {"type": "safety_adjusted", "rationale": "已做安全调整"}
        with self.assertRaisesRegex(ValueError, "必须修改|仍不安全"):
            build_docx.validate_handoff(valid_handoff([item]))

    def test_unknown_public_keys_and_locks_are_rejected(self):
        public = {"filename": "test", "items": [valid_item()["public"]]}
        public["items"][0]["locks"] = {}
        with self.assertRaises(ValueError):
            build_docx.validate_public_input(public)

    def test_title_decision_is_enforced(self):
        data = valid_handoff()
        data["items"][0]["public"]["title"] = "改过的标题"
        with self.assertRaisesRegex(ValueError, "unchanged"):
            build_docx.validate_handoff(data)
        data["items"][0]["title_decision"] = {"type": "safety_adjusted", "rationale": "去掉危险医学确定性，保留数字悬念"}
        build_docx.validate_handoff(data)

    def test_filename_boundaries(self):
        for filename in ("../x", "folder\\x", "C:/x", "CON", "CON.txt", "PRN.anything"):
            with self.subTest(filename=filename):
                with self.assertRaises(ValueError):
                    build_docx.safe_filename(filename)

    def test_existing_output_is_not_overwritten(self):
        data = valid_handoff(filename="existing")
        with tempfile.TemporaryDirectory() as output_dir:
            existing = Path(output_dir) / "existing.docx"
            existing.write_bytes(b"keep")
            with self.assertRaises(FileExistsError):
                build_docx.build(data, output_dir)
            self.assertEqual(b"keep", existing.read_bytes())

    def test_docx_layout_single_and_multi(self):
        with tempfile.TemporaryDirectory() as output_dir:
            single_path, _, _ = build_docx.build(valid_handoff(filename="single"), output_dir)
            single = Document(single_path)
            self.assertFalse(any(p.text.startswith("第1条") for p in single.paragraphs))
            title = next(p for p in single.paragraphs if p.text.startswith("标题："))
            self.assertTrue(all(run.bold for run in title.runs))
            body = next(p for p in single.paragraphs if "桃树叶一大把" in p.text)
            self.assertEqual("宋体", body.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")))

            multi_path, _, _ = build_docx.build(valid_handoff([valid_item(), valid_item()], filename="multi"), output_dir)
            multi = Document(multi_path)
            self.assertIn("第1条", [p.text for p in multi.paragraphs])
            self.assertIn("第2条", [p.text for p in multi.paragraphs])
            self.assertTrue(any(p._p.xpath("./w:pPr/w:pBdr/w:bottom") for p in multi.paragraphs))


class BuildDocxCliE2ETests(unittest.TestCase):
    script = str(SCRIPTS / "build_docx.py")

    def run_cli(self, handoff, handoff_dir, output_dir):
        handoff_path = Path(handoff_dir) / "handoff.json"
        handoff_path.write_text(json.dumps(handoff, ensure_ascii=False), encoding="utf-8")
        return subprocess.run(
            [sys.executable, self.script, str(handoff_path), str(output_dir)],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

    def test_approved_gallstone_v3_fixture_exports_through_public_cli(self):
        fixture = ROOT / "tests" / "fixtures" / "02-gallstones"
        data = json.loads((fixture / "approved-handoff-v3.json").read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            output_path = Path(output_dir) / "胆结石-传统洗稿回归.docx"
            self.assertTrue(output_path.is_file())
            paragraphs = [p.text for p in Document(output_path).paragraphs]
            expected = build_docx.split_script(data["items"][0]["public"]["script"])
            body_start = next(i for i, text in enumerate(paragraphs) if text.startswith("标题：")) + 1
            self.assertEqual(expected, paragraphs[body_start:body_start + len(expected)])
            with zipfile.ZipFile(output_path) as archive:
                xml_text = archive.read("word/document.xml").decode("utf-8")
            for forbidden in (
                "hooks", "review_record", "warning_decisions", "provenance",
                "approval_status", "safety_disposition", "source_span",
            ):
                self.assertNotIn(forbidden, xml_text)

    def test_v3_rewritten_handoff_exports_through_public_cli(self):
        source_text = "眼睛模糊看不清，方中用青葙子10到15克，证型为肝火上炎。"
        formula_start = source_text.index("青葙子10到15克")
        disease_start = source_text.index("眼睛模糊看不清")
        syndrome_start = source_text.index("肝火上炎")

        def hook(hook_id, kind, raw, start):
            return {
                "id": hook_id, "kind": kind, "raw": raw, "delivery_text": raw,
                "source_span": [start, start + len(raw)], "normalization_ops": [],
                "provenance": "source_verbatim", "approval_status": "approved",
                "safety_disposition": "verbatim",
            }

        script = (
            "眼睛模糊看不清，先别急着花冤枉钱。\n"
            "在中医看来，这类表现可能和肝火上扰有关。大家可以取青葙子10到15克，煎水服用，有助于清泄肝火、改善目赤翳障（需辨证为肝火上炎）。\n"
            "记不住先收藏，用到时再翻出来看。"
        )
        findings = [{
            "id": "finding-1", "kind": "contraindication",
            "text": "禁忌人群，青光眼患者不建议使用。", "target": "tips",
            "approval_status": "approved", "source": "",
        }]
        data = {
            "schema_version": "3", "stage": "rewritten", "filename": "v3-cli-round-trip",
            "items": [{
                "source": {"link": "", "title": "眼睛模糊看不清，教你一个土方法", "text": source_text},
                "hooks": {"verbatim": [
                    hook("disease-1", "disease", "眼睛模糊看不清", disease_start),
                    hook("formula-1", "formula_item", "青葙子10到15克", formula_start),
                    hook("syndrome-1", "syndrome", "肝火上炎", syndrome_start),
                ], "semantic_strength": []},
                "review_findings": findings,
                "title_review": {"decision": "unchanged", "title": "眼睛模糊看不清，教你一个土方法", "reason_codes": [], "rationale": ""},
                "review_record": {key: True for key in build_docx.MANUAL_CHECK_KEYS},
                "warning_decisions": {},
                "public": {
                    "link": "", "title": "眼睛模糊看不清，教你一个土方法", "script": script,
                    "notes": [], "tips": [{"text": findings[0]["text"]}], "processing": [],
                },
            }],
        }
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            output_path = Path(output_dir) / "v3-cli-round-trip.docx"
            self.assertTrue(output_path.is_file())
            text = "\n".join(p.text for p in Document(output_path).paragraphs)
            self.assertIn("青葙子10到15克", text)
            self.assertIn("青光眼患者不建议使用", text)

    def test_v3_pending_hook_and_finding_block_export(self):
        source_text = "原文含一项内容。"
        base_item = {
            "source": {"link": "", "title": "测试标题", "text": source_text},
            "hooks": {"verbatim": [], "semantic_strength": []},
            "review_findings": [{
                "id": "finding-1", "kind": "contraindication", "text": "待批准提醒。",
                "target": "tips", "approval_status": "pending", "source": "",
            }],
            "title_review": {"decision": "unchanged", "title": "测试标题", "reason_codes": [], "rationale": ""},
            "review_record": {key: True for key in build_docx.MANUAL_CHECK_KEYS},
            "warning_decisions": {},
            "public": {"link": "", "title": "测试标题", "script": "普通合规正文。", "notes": [], "tips": [], "processing": []},
        }
        for mutation in ("finding", "hook"):
            item = json.loads(json.dumps(base_item, ensure_ascii=False))
            if mutation == "hook":
                item["review_findings"] = []
                item["hooks"]["verbatim"] = [{
                    "id": "formula-1", "kind": "formula_item", "raw": "原文",
                    "delivery_text": "原文", "source_span": [0, 2], "normalization_ops": [],
                    "provenance": "source_verbatim", "approval_status": "pending",
                    "safety_disposition": "pending",
                }]
            data = {"schema_version": "3", "stage": "rewritten", "filename": "v3-pending-" + mutation, "items": [item]}
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
                result = self.run_cli(data, handoff_dir, output_dir)
                self.assertNotEqual(0, result.returncode)
                self.assertEqual([], list(Path(output_dir).iterdir()))

    def test_v3_fixed_phrase_is_enforced(self):
        source_text = "固定口诀必须牢记。"
        data = {
            "schema_version": "3", "stage": "rewritten", "filename": "v3-fixed-phrase",
            "items": [{
                "source": {"link": "", "title": "测试标题", "text": source_text},
                "hooks": {"verbatim": [{
                    "id": "fixed-1", "kind": "fixed_phrase", "raw": "固定口诀必须牢记",
                    "delivery_text": "固定口诀必须牢记", "source_span": [0, 8],
                    "normalization_ops": [], "provenance": "source_verbatim",
                    "approval_status": "approved", "safety_disposition": "verbatim",
                }], "semantic_strength": []},
                "review_findings": [],
                "title_review": {"decision": "unchanged", "title": "测试标题", "reason_codes": [], "rationale": ""},
                "review_record": {key: True for key in build_docx.MANUAL_CHECK_KEYS},
                "warning_decisions": {},
                "public": {"link": "", "title": "测试标题", "script": "普通合规正文。", "notes": [], "tips": [], "processing": []},
            }],
        }
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertNotEqual(0, result.returncode)
            self.assertEqual([], list(Path(output_dir).iterdir()))

    def test_json_file_to_docx_cli_round_trip(self):
        data = valid_handoff(filename="cli-round-trip")
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            output_path = Path(output_dir) / "cli-round-trip.docx"
            self.assertTrue(output_path.is_file())
            text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
            self.assertIn(data["items"][0]["public"]["script"], text)
            self.assertNotIn("manual_checks", text)
            self.assertNotIn("locks", text)
            with zipfile.ZipFile(output_path) as archive:
                xml_text = "\n".join(
                    archive.read(name).decode("utf-8")
                    for name in archive.namelist()
                    if name.endswith(".xml")
                )
            for forbidden in ("locks", "manual_checks", "warning_decisions", "provenance", "internal_review"):
                self.assertNotIn(forbidden, xml_text)

    def test_cli_accepts_utf8_bom_handoff(self):
        data = valid_handoff(filename="bom-round-trip")
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            handoff_path = Path(handoff_dir) / "handoff.json"
            handoff_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8-sig")
            result = subprocess.run(
                [sys.executable, self.script, str(handoff_path), str(output_dir)],
                capture_output=True,
            )
            result.stdout.decode("utf-8", errors="strict")
            result.stderr.decode("utf-8", errors="strict")
            self.assertEqual(0, result.returncode)

    def test_cli_rejects_existing_output_without_changing_bytes(self):
        data = valid_handoff(filename="existing-cli")
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            existing = Path(output_dir) / "existing-cli.docx"
            existing.write_bytes(b"keep")
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertNotEqual(0, result.returncode)
            self.assertEqual(b"keep", existing.read_bytes())

    def test_invalid_item_keeps_batch_output_atomic(self):
        data = valid_handoff([valid_item(), valid_item(script="")], filename="invalid-batch")
        with tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
            result = self.run_cli(data, handoff_dir, output_dir)
            self.assertNotEqual(0, result.returncode)
            self.assertIn("导出失败", result.stderr)
            self.assertEqual([], list(Path(output_dir).iterdir()))

    def test_editorial_tone_and_invalid_syndrome_keep_output_atomic(self):
        cases = (
            valid_item(script="腰腿疼痛关键是先看自己到底适不适合，可以准备桃树叶一大把、生姜五片（需辨证为寒湿痹阻）。"),
            valid_item(script="腰腿疼痛可以准备桃树叶一大把、生姜五片（需辨证为寒湿痹阻）（需辨证为寒湿痹阻）。"),
        )
        for index, invalid_item in enumerate(cases):
            with self.subTest(index=index), tempfile.TemporaryDirectory() as handoff_dir, tempfile.TemporaryDirectory() as output_dir:
                data = valid_handoff([valid_item(), invalid_item], filename="invalid-semantic-batch")
                result = self.run_cli(data, handoff_dir, output_dir)
                self.assertNotEqual(0, result.returncode)
                self.assertIn("导出失败", result.stderr)
                self.assertEqual([], list(Path(output_dir).iterdir()))


class FixtureContractTests(unittest.TestCase):
    fixtures_dir = ROOT / "tests" / "fixtures"

    def test_fixture_contracts_are_complete_and_consistent(self):
        case_paths = sorted(self.fixtures_dir.glob("*/case.json"))
        self.assertGreaterEqual(len(case_paths), 3)
        required_keys = {"id", "source", "locks", "machine_assertions", "manual_dimensions"}
        for path in case_paths:
            with self.subTest(path=path):
                case = json.loads(path.read_text(encoding="utf-8"))
                self.assertEqual(required_keys, set(case))
                self.assertEqual(path.parent.name, case["id"])
                self.assertEqual({"link", "title", "text"}, set(case["source"]))
                self.assertEqual(build_docx.LOCK_KEYS, set(case["locks"]))
                for key in build_docx.LOCK_KEYS:
                    build_docx._validate_lock_list(case["locks"][key], "locks." + key)
                assertions = case["machine_assertions"]
                self.assertEqual(
                    {"required_literals", "forbidden_literals", "max_chars", "title_decision"},
                    set(assertions),
                )
                self.assertEqual(300, assertions["max_chars"])
                self.assertIn(assertions["title_decision"], {"unchanged", "safety_adjusted"})
                for key in ("required_literals", "forbidden_literals"):
                    self.assertTrue(all(isinstance(value, str) and value for value in assertions[key]))
                self.assertTrue(all(isinstance(value, str) and value for value in case["manual_dimensions"]))
                self.assertTrue((path.parent / "input.txt").is_file())
                self.assertTrue((path.parent / "expected_checks.md").is_file())
                self.assertEqual(case["source"]["text"].strip(), (path.parent / "input.txt").read_text(encoding="utf-8").strip())

    def test_approved_fixture_outputs_drive_validator_and_handoff(self):
        for approved_path in sorted(self.fixtures_dir.glob("*/approved-output.txt")):
            case_path = approved_path.parent / "case.json"
            case = json.loads(case_path.read_text(encoding="utf-8"))
            script = approved_path.read_text(encoding="utf-8").strip()
            assertions = case["machine_assertions"]
            for literal in assertions["required_literals"]:
                self.assertIn(literal, script, approved_path.parent.name)
            for literal in assertions["forbidden_literals"]:
                self.assertNotIn(literal, script, approved_path.parent.name)
            report = validate_output.validate_text(script, case["locks"], case["source"]["text"])
            self.assertEqual([], report["errors"], approved_path.parent.name)
            self.assertLessEqual(report["char_count"], assertions["max_chars"])
            item = valid_item(script=script, source_text=case["source"]["text"])
            item["source"] = case["source"]
            item["locks"] = case["locks"]
            item["public"]["link"] = case["source"]["link"]
            item["public"]["title"] = case["source"]["title"]
            item["title_decision"] = {"type": assertions["title_decision"], "rationale": ""}
            item["warning_decisions"] = {}
            build_docx.validate_handoff(valid_handoff([item], filename=approved_path.parent.name))

    def test_gallstone_confirmed_and_defensive_outputs_are_frozen(self):
        fixture = self.fixtures_dir / "02-gallstones"
        approved = (fixture / "approved-output.txt").read_text(encoding="utf-8")
        failed = (fixture / "failed-output-defensive-popsci.txt").read_text(encoding="utf-8")
        self.assertIn("建议先送上一颗小爱心，再留下一句谢谢", approved)
        self.assertIn("可以取海金沙9克、金钱草15克、鸡内金9克", approved)
        approved_ids = {x["id"] for x in validate_output.validate_text(approved)["errors"]}
        self.assertFalse({"popsci-bridge", "editorial-tone"}.intersection(approved_ids))
        self.assertIn("常见的配伍思路", failed)
        failed_ids = {x["id"] for x in validate_output.validate_text(failed)["errors"]}
        self.assertIn("popsci-bridge", failed_ids)
        self.assertIn("editorial-tone", failed_ids)

    def test_blurred_vision_complete_failure_outputs_are_frozen(self):
        fixture = self.fixtures_dir / "03-blurred-vision"
        repeated = (fixture / "failed-output-repeated-preview.txt").read_text(encoding="utf-8")
        popsci = (fixture / "failed-output-popsci-bridge.txt").read_text(encoding="utf-8")
        editorial = (fixture / "failed-output-editorial-tone.txt").read_text(encoding="utf-8")
        approved = (fixture / "approved-output.txt").read_text(encoding="utf-8")
        self.assertIn("今天再把这个土方法讲透", repeated)
        self.assertIn("today-preview-repeat", {x["id"] for x in validate_output.validate_text(repeated)["errors"]})
        self.assertIn("顺着这个思路，常会用到", popsci)
        self.assertIn("popsci-bridge", {x["id"] for x in validate_output.validate_text(popsci)["errors"]})
        self.assertIn("关键是先看自己到底适不适合", editorial)
        self.assertIn("editorial-tone", {x["id"] for x in validate_output.validate_text(editorial)["errors"]})
        self.assertIn("大家可以取青葙子10到15克", approved)
        self.assertNotIn("today-preview-repeat", {x["id"] for x in validate_output.validate_text(approved)["errors"]})
        self.assertNotIn("popsci-bridge", {x["id"] for x in validate_output.validate_text(approved)["errors"]})

    def test_blurred_vision_fixture_freezes_real_failure_mode(self):
        path = self.fixtures_dir / "03-blurred-vision" / "case.json"
        case = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual("眼睛模糊看不清，教你一个土方法", case["source"]["title"])
        self.assertEqual(["肝火上炎"], validate_output.lock_texts(case["locks"]["syndromes"]))
        self.assertEqual(["青葙子10到15克"], validate_output.lock_texts(case["locks"]["formula_items"]))
        assertions = case["machine_assertions"]
        self.assertIn("（需辨证为肝火上炎）", assertions["required_literals"])
        for text in ("关键是先看自己到底适不适合", "视物模糊不一定都是这个原因", "在辨证合适的前提下"):
            self.assertIn(text, assertions["forbidden_literals"])
            self.assertIn("editorial-tone", {x["id"] for x in validate_output.validate_text(text)["errors"]})


if __name__ == "__main__":
    unittest.main()

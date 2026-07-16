"""医誉康达口播 canonical handoff 校验与 Word 渲染。"""
import json
import os
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))
from handoff_contract import (
    approved_findings_projection, compile_locks, validate_v3_handoff_shape,
)
from validate_output import find_medical_safety, validate_text

BODY_FONT = "宋体"
BODY_SIZE = 10.5
HEADING_SIZE = 10.5
SECTION_ORDER = ("notes", "tips", "processing")
HANDOFF_KEYS = {"schema_version", "stage", "filename", "items"}
HANDOFF_ITEM_KEYS = {"source", "locks", "title_decision", "manual_checks", "warning_decisions", "public"}
SOURCE_KEYS = {"link", "title", "text"}
LOCK_KEYS = {"formula_items", "diseases", "syndromes", "numeric_hooks"}
LOCK_ITEM_KEYS = {"raw", "locked_text", "normalization", "provenance", "requires_approval"}
TITLE_DECISION_KEYS = {"type", "rationale"}
MANUAL_CHECK_KEYS = {
    "narrative_chain", "hook_and_title", "rewrite_quality", "medical_boundary",
    "lock_provenance", "oral_naturalness", "guidance_fit",
}
PUBLIC_TOP_KEYS = {"filename", "items"}
PUBLIC_ITEM_KEYS = {"link", "title", "script", "notes", "tips", "processing"}
ENTRY_KEYS = {"text", "source"}
WARNING_STATUSES = {"false_positive", "accepted_with_reason"}
FORBIDDEN_PUBLIC_PHRASES = (
    "以上为通用知识，请执业医师复核", "review_note",
    "licensed_physician_review_required",
)
WINDOWS_RESERVED = {"CON", "PRN", "AUX", "NUL", *("COM%d" % i for i in range(1, 10)), *("LPT%d" % i for i in range(1, 10))}


def _exact_keys(value, allowed, path):
    if not isinstance(value, dict):
        raise ValueError("%s 必须是对象" % path)
    unknown = sorted(set(value) - allowed)
    missing = sorted(allowed - set(value))
    if unknown:
        raise ValueError("%s 含未知字段：%s" % (path, ",".join(unknown)))
    if missing:
        raise ValueError("%s 缺字段：%s" % (path, ",".join(missing)))


def _string(value, path, allow_empty=False):
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError("%s 必须是%s字符串" % (path, "" if allow_empty else "非空"))
    return value.strip()


def _validate_lock_list(values, path):
    if not isinstance(values, list):
        raise ValueError("%s 必须是数组" % path)
    for index, value in enumerate(values):
        item_path = "%s[%d]" % (path, index)
        if isinstance(value, str):
            _string(value, item_path)
            continue
        _exact_keys(value, LOCK_ITEM_KEYS, item_path)
        _string(value["raw"], item_path + ".raw")
        _string(value["locked_text"], item_path + ".locked_text")
        if not isinstance(value["normalization"], list) or not all(isinstance(x, str) for x in value["normalization"]):
            raise ValueError(item_path + ".normalization 必须是字符串数组")
        if value["provenance"] not in {"source_verbatim", "source_normalized", "user_required", "review_approved"}:
            raise ValueError(item_path + ".provenance 无效")
        if not isinstance(value["requires_approval"], bool):
            raise ValueError(item_path + ".requires_approval 必须是布尔值")
        if value["requires_approval"]:
            raise ValueError(item_path + " 尚未批准，不能导出")


def _validate_entries(entries, path):
    if not isinstance(entries, list):
        raise ValueError(path + " 必须是数组")
    for index, entry in enumerate(entries):
        item_path = "%s[%d]" % (path, index)
        if not isinstance(entry, dict):
            raise ValueError(item_path + " 必须是对象")
        unknown = sorted(set(entry) - ENTRY_KEYS)
        if unknown:
            raise ValueError("%s 含未知字段：%s" % (item_path, ",".join(unknown)))
        if "text" not in entry:
            raise ValueError(item_path + " 缺 text")
        _string(entry["text"], item_path + ".text")
        if "source" in entry and entry["source"] is not None:
            _string(entry["source"], item_path + ".source", allow_empty=True)


def validate_public_input(data):
    """严格验证公开投影；未知字段和内部数据 fail closed。"""
    _exact_keys(data, PUBLIC_TOP_KEYS, "public")
    _string(data["filename"], "public.filename")
    if not isinstance(data["items"], list) or not data["items"]:
        raise ValueError("public.items 必须是非空数组")
    for index, item in enumerate(data["items"]):
        path = "public.items[%d]" % index
        _exact_keys(item, PUBLIC_ITEM_KEYS, path)
        _string(item["link"], path + ".link", allow_empty=True)
        _string(item["title"], path + ".title")
        _string(item["script"], path + ".script")
        for key in SECTION_ORDER:
            _validate_entries(item[key], path + "." + key)
        serialized = json.dumps(item, ensure_ascii=False)
        for phrase in FORBIDDEN_PUBLIC_PHRASES:
            if phrase in serialized:
                raise ValueError("%s 含禁用内部文本 %s" % (path, phrase))


def _warning_decision(decisions, issue_id):
    value = decisions.get(issue_id)
    if isinstance(value, str):
        return value, ""
    if isinstance(value, dict):
        unknown = set(value) - {"status", "reason"}
        if unknown:
            return None, ""
        return value.get("status"), value.get("reason", "")
    return None, ""


def _validate_public_medical_safety(public, path):
    title_issues = find_medical_safety(public["title"], path + ".title")
    if title_issues:
        raise ValueError("%s 标题医疗安全校验失败：%s" % (path, "；".join(x["message"] for x in title_issues)))
    for key in SECTION_ORDER:
        _validate_entries(public[key], path + "." + key)
        for entry_index, entry in enumerate(public[key]):
            issues = find_medical_safety(entry["text"], "%s.%s[%d].text" % (path, key, entry_index))
            if issues:
                raise ValueError("%s %s 医疗安全校验失败：%s" % (
                    path, key, "；".join(x["message"] for x in issues),
                ))


def _validate_warnings(decisions, report, path):
    for warning in report["warnings"]:
        status, reason = _warning_decision(decisions, warning["id"])
        if status not in WARNING_STATUSES:
            raise ValueError("%s warning [%s] 未处置或状态无效" % (path, warning["id"]))
        if not isinstance(reason, str) or not reason.strip():
            raise ValueError("%s warning [%s] 必须填写理由" % (path, warning["id"]))
    unknown = sorted(set(decisions) - {x["id"] for x in report["warnings"]})
    if unknown:
        raise ValueError("%s warning_decisions 含未知或已过期项：%s" % (path, ",".join(unknown)))


def validate_v3_handoff(data):
    """Validate v3 rewritten handoff and compile only approved constraints/findings."""
    validate_v3_handoff_shape(data)
    public_items = []
    reports = []
    for index, item in enumerate(data["items"]):
        path = "items[%d]" % index
        public = item["public"]
        _exact_keys(public, PUBLIC_ITEM_KEYS, path + ".public")
        if public["link"] != item["source"]["link"]:
            raise ValueError(path + " public.link 必须与 source.link 一致")
        findings = approved_findings_projection(item["review_findings"])
        for key in SECTION_ORDER:
            if public[key] != findings[key]:
                raise ValueError("%s public.%s 必须由 approved review_findings 确定性投影" % (path, key))
        _validate_public_medical_safety(public, path + ".public")
        locks = compile_locks(item["hooks"])
        report = validate_text(public["script"], locks, item["source"]["text"])
        if report["errors"]:
            raise ValueError("%s 正文校验失败：%s" % (path, "；".join(x["message"] for x in report["errors"])))
        _validate_warnings(item["warning_decisions"], report, path)
        public_items.append(public)
        reports.append(report)
    public_data = {"filename": data["filename"], "items": public_items}
    validate_public_input(public_data)
    return public_data, reports


def validate_handoff(data):
    version = data.get("schema_version") if isinstance(data, dict) else None
    if version == "3":
        return validate_v3_handoff(data)
    if version != "2":
        raise ValueError("schema_version 必须为 2 或 3")
    return validate_v2_handoff(data)


def validate_v2_handoff(data):
    _exact_keys(data, HANDOFF_KEYS, "handoff")
    if data["schema_version"] != "2":
        raise ValueError("schema_version 必须为 2")
    if data["stage"] != "rewritten":
        raise ValueError("stage 必须为 rewritten，validated/exported 只能由脚本产生")
    _string(data["filename"], "filename")
    if not isinstance(data["items"], list) or not data["items"]:
        raise ValueError("items 必须是非空数组")

    public_items = []
    reports = []
    for index, item in enumerate(data["items"]):
        path = "items[%d]" % index
        _exact_keys(item, HANDOFF_ITEM_KEYS, path)
        _exact_keys(item["source"], SOURCE_KEYS, path + ".source")
        source = item["source"]
        _string(source["link"], path + ".source.link", allow_empty=True)
        source_title = _string(source["title"], path + ".source.title")
        source_text = _string(source["text"], path + ".source.text")

        _exact_keys(item["locks"], LOCK_KEYS, path + ".locks")
        for key in LOCK_KEYS:
            _validate_lock_list(item["locks"][key], path + ".locks." + key)

        _exact_keys(item["title_decision"], TITLE_DECISION_KEYS, path + ".title_decision")
        decision_type = item["title_decision"]["type"]
        rationale = item["title_decision"]["rationale"]
        if decision_type not in {"unchanged", "safety_adjusted"}:
            raise ValueError(path + ".title_decision.type 无效")
        if not isinstance(rationale, str):
            raise ValueError(path + ".title_decision.rationale 必须是字符串")

        _exact_keys(item["manual_checks"], MANUAL_CHECK_KEYS, path + ".manual_checks")
        failed_checks = [key for key, passed in item["manual_checks"].items() if passed is not True]
        if failed_checks:
            raise ValueError("%s 人工卡口未通过：%s" % (path, ",".join(failed_checks)))
        if not isinstance(item["warning_decisions"], dict):
            raise ValueError(path + ".warning_decisions 必须是对象")

        public = item["public"]
        _exact_keys(public, PUBLIC_ITEM_KEYS, path + ".public")
        public_title = _string(public["title"], path + ".public.title")
        if decision_type == "unchanged" and public_title != source_title:
            raise ValueError(path + " 标题声明 unchanged，但 public.title 与 source.title 不一致")
        if decision_type == "safety_adjusted":
            if not rationale.strip():
                raise ValueError(path + " safety_adjusted 必须填写 rationale")
            if public_title == source_title:
                raise ValueError(path + " safety_adjusted 必须修改标题")
        title_issues = find_medical_safety(public_title, path + ".public.title")
        if title_issues:
            raise ValueError("%s 标题医疗安全校验失败：%s" % (path, "；".join(x["message"] for x in title_issues)))
        if public["link"] != source["link"]:
            raise ValueError(path + " public.link 必须与 source.link 一致")
        for key in SECTION_ORDER:
            _validate_entries(public[key], path + ".public." + key)
            for entry_index, entry in enumerate(public[key]):
                section_issues = find_medical_safety(
                    entry["text"], "%s.public.%s[%d].text" % (path, key, entry_index)
                )
                if section_issues:
                    raise ValueError("%s %s 医疗安全校验失败：%s" % (
                        path, key, "；".join(x["message"] for x in section_issues),
                    ))

        report = validate_text(public["script"], item["locks"], source_text)
        if report["errors"]:
            raise ValueError("%s 正文校验失败：%s" % (path, "；".join(x["message"] for x in report["errors"])))
        for warning in report["warnings"]:
            status, reason = _warning_decision(item["warning_decisions"], warning["id"])
            if status not in WARNING_STATUSES:
                raise ValueError("%s warning [%s] 未处置或状态无效" % (path, warning["id"]))
            if not isinstance(reason, str) or not reason.strip():
                raise ValueError("%s warning [%s] 必须填写理由" % (path, warning["id"]))
        unknown_decisions = sorted(set(item["warning_decisions"]) - {x["id"] for x in report["warnings"]})
        if unknown_decisions:
            raise ValueError("%s warning_decisions 含未知或已过期项：%s" % (path, ",".join(unknown_decisions)))

        public_items.append(public)
        reports.append(report)

    public_data = {"filename": data["filename"], "items": public_items}
    validate_public_input(public_data)
    return public_data, reports


def safe_filename(value):
    filename = _string(value, "filename")
    if filename.lower().endswith(".docx"):
        filename = filename[:-5]
    if filename in {".", ".."} or "/" in filename or "\\" in filename:
        raise ValueError("filename 必须是安全 basename，不能含路径")
    if re.search(r'[<>:"/\\|?*\x00-\x1f]', filename) or filename.endswith((" ", ".")):
        raise ValueError("filename 含 Windows 非法字符")
    device_stem = filename.split(".", 1)[0].rstrip(" .").upper()
    if device_stem in WINDOWS_RESERVED:
        raise ValueError("filename 是 Windows 保留名")
    final_name = filename + ".docx"
    if len(final_name) > 240:
        raise ValueError("filename 过长")
    return final_name


def set_cjk_font(run, name=BODY_FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_inline_label(doc, label, content, bold_content=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.size = Pt(HEADING_SIZE)
    set_cjk_font(label_run)
    content_run = paragraph.add_run(content)
    content_run.font.size = Pt(BODY_SIZE)
    content_run.bold = bold_content
    set_cjk_font(content_run)


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(BODY_SIZE)
    set_cjk_font(run)


def add_separator(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(10)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BFBFBF")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def split_script(script):
    lines = script.split("\n")
    if any(line != line.strip() for line in lines):
        raise ValueError("正文每行不得含首尾空白")
    paragraphs = [line for line in lines if line]
    if not paragraphs:
        raise ValueError("正文不能为空")
    return paragraphs


def split_tips(text):
    combined = text.strip()
    marker = "禁忌人群"
    if marker not in combined:
        return [combined] if combined else []
    index = combined.index(marker)
    return [part for part in (combined[:index].strip(), combined[index:].strip()) if part]


def render_item(doc, item):
    if item["link"]:
        add_inline_label(doc, "视频链接：", item["link"])
    add_inline_label(doc, "标题：", item["title"], bold_content=True)
    for paragraph in split_script(item["script"]):
        add_body(doc, paragraph)
    for key in SECTION_ORDER:
        entries = item[key]
        if not entries:
            continue
        combined = "\n".join(entry["text"].strip() for entry in entries)
        parts = split_tips(combined) if key == "tips" else [combined]
        for part in parts:
            add_body(doc, "【%s】" % part)


def build(data, output_dir):
    public_data, reports = validate_handoff(data)
    output_root = Path(output_dir).resolve()
    if not output_root.is_dir():
        raise ValueError("输出目录不存在：%s" % output_root)
    filename = safe_filename(public_data["filename"])
    output_path = (output_root / filename).resolve()
    if output_path.parent != output_root:
        raise ValueError("输出路径越界")
    if output_path.exists():
        raise FileExistsError("输出文件已存在，拒绝覆盖：%s" % output_path)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    multi = len(public_data["items"]) > 1
    for index, item in enumerate(public_data["items"]):
        if index:
            add_separator(doc)
        if multi:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run("第%d条" % (index + 1))
            run.bold = True
            run.font.size = Pt(12)
            set_cjk_font(run)
        render_item(doc, item)

    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix="yiyukangda-", suffix=".docx", dir=output_root, delete=False) as handle:
            temp_path = Path(handle.name)
        doc.save(temp_path)
        # 原子发布且目标必须不存在；避免 exists()+replace 的并发覆盖窗口。
        try:
            os.link(temp_path, output_path)
        except FileExistsError:
            raise FileExistsError("输出文件在生成期间出现，拒绝覆盖：%s" % output_path)
        except OSError as exc:
            raise OSError("当前文件系统不支持安全无覆盖发布：%s" % exc)
        temp_path.unlink()
        temp_path = None
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink()
    return str(output_path), len(public_data["items"]), reports


def main():
    if len(sys.argv) < 2:
        raise SystemExit("用法：python build_docx.py <canonical-handoff.json> [output_dir]")
    input_path = Path(sys.argv[1])
    output_dir = sys.argv[2] if len(sys.argv) > 2 else str(Path.home() / "Desktop")
    try:
        data = json.loads(input_path.read_text(encoding="utf-8-sig"))
        output_path, count, _reports = build(data, output_dir)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError, FileExistsError) as exc:
        raise SystemExit("导出失败：%s" % exc)
    print("已校验并生成 %d 条 → %s" % (count, output_path))


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
